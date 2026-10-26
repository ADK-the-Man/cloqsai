"""
Excel (.xlsx) workbook generator using openpyxl with multiple sheets, formatting, and formulas.
"""
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime, timezone
from typing import Any, Dict, List
from core.storage import save_artifact


async def generate_estimate_xlsx(
    project_name: str,
    provider: str,
    scenario_name: str,
    line_items: List[Dict[str, Any]],
    assumptions: List[str],
    exclusions: List[str],
    user_id: str,
    job_id: str,
    provenance_label: str,
    estimate_url: str = None
) -> tuple[str, int, str]:
    wb = openpyxl.Workbook()
    
    # ── Sheet 1: Summary ──────────────────────────────────────────────
    ws_summary = wb.active
    ws_summary.title = "Executive Summary"
    ws_summary.views.sheetView[0].showGridLines = True

    header_font = Font(name="Segoe UI", size=16, bold=True, color="1E3A8A")
    title_cell = ws_summary["B2"]
    title_cell.value = "CLOQS AI — AI-Powered Cloud Architecture & Cost Intelligence Plan"
    title_cell.font = header_font

    meta_font = Font(name="Segoe UI", size=11, bold=True)
    val_font = Font(name="Segoe UI", size=11)

    metadata = [
        ("Project Name", project_name),
        ("Cloud Provider", provider.upper()),
        ("Architecture Scenario", scenario_name),
        ("Official Calculator Direct Link", estimate_url or "https://calculator.aws/#/estimate" if provider.lower() == "aws" else "https://azure.microsoft.com/en-us/pricing/calculator/"),
        ("Export Timestamp (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")),
        ("Data Source Provenance", provenance_label),
    ]

    for idx, (label, val) in enumerate(metadata, start=4):
        ws_summary.cell(row=idx, column=2, value=label).font = meta_font
        ws_summary.cell(row=idx, column=3, value=val).font = val_font

    total_monthly = sum(float(x.get("monthly_cost_usd", 0.0)) for x in line_items)
    total_annual = total_monthly * 12

    ws_summary.cell(row=11, column=2, value="Total Estimated Monthly Spend:").font = Font(name="Segoe UI", size=13, bold=True, color="047857")
    ws_summary.cell(row=11, column=3, value=f"${total_monthly:,.2f} USD").font = Font(name="Segoe UI", size=13, bold=True, color="047857")
    ws_summary.cell(row=12, column=2, value="Total Estimated Annual Spend:").font = Font(name="Segoe UI", size=13, bold=True, color="047857")
    ws_summary.cell(row=12, column=3, value=f"${total_annual:,.2f} USD").font = Font(name="Segoe UI", size=13, bold=True, color="047857")

    # ── Sheet 2: Resource Breakdown ───────────────────────────────────
    ws_items = wb.create_sheet(title="Resource Line Items")
    ws_items.views.sheetView[0].showGridLines = True

    columns = [
        "Resource Name", "Service Name", "Region", "SKU / Size",
        "Quantity", "Availability Tier", "Pricing Model",
        "Monthly Cost (USD)", "Annual Cost (USD)", "Pricing Source", "Assumptions"
    ]

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font_white = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

    for col_num, col_title in enumerate(columns, 1):
        cell = ws_items.cell(row=1, column=col_num, value=col_title)
        cell.fill = header_fill
        cell.font = header_font_white
        cell.alignment = Alignment(horizontal="center", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )

    row_num = 2
    for item in line_items:
        monthly = float(item.get("monthly_cost_usd", 0.0))
        annual = float(item.get("annual_cost_usd", monthly * 12))
        item_assumptions = "; ".join(item.get("assumptions", [])) if isinstance(item.get("assumptions"), list) else str(item.get("assumptions", ""))

        row_vals = [
            item.get("resource_name", ""),
            item.get("service_name", ""),
            item.get("region", "us-east-1" if provider.lower() == "aws" else "eastus"),
            item.get("sku", ""),
            item.get("quantity", 1),
            item.get("availability_config", "Standard"),
            item.get("pricing_model", "On-Demand"),
            monthly,
            annual,
            item.get("pricing_source", provenance_label),
            item_assumptions
        ]

        for col_num, val in enumerate(row_vals, 1):
            c = ws_items.cell(row=row_num, column=col_num, value=val)
            c.font = Font(name="Segoe UI", size=10)
            c.border = thin_border
            if col_num in (8, 9):
                c.number_format = "$#,##0.00"
                c.alignment = Alignment(horizontal="right")
            elif col_num == 5:
                c.alignment = Alignment(horizontal="center")

        row_num += 1

    # Formula row
    total_cell_label = ws_items.cell(row=row_num, column=1, value="TOTAL")
    total_cell_label.font = Font(name="Segoe UI", size=11, bold=True)
    
    cost_col_8 = ws_items.cell(row=row_num, column=8, value=f"=SUM(H2:H{row_num-1})")
    cost_col_8.font = Font(name="Segoe UI", size=11, bold=True)
    cost_col_8.number_format = "$#,##0.00"

    cost_col_9 = ws_items.cell(row=row_num, column=9, value=f"=SUM(I2:I{row_num-1})")
    cost_col_9.font = Font(name="Segoe UI", size=11, bold=True)
    cost_col_9.number_format = "$#,##0.00"

    # Auto-adjust column widths
    for col in ws_items.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_items.column_dimensions[col_letter].width = max(max_len + 3, 12)

    for col in ws_summary.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_summary.column_dimensions[col_letter].width = max(max_len + 3, 15)

    # ── Sheet 3: Assumptions & Exclusions ─────────────────────────────
    ws_notes = wb.create_sheet(title="Assumptions & Scope")
    ws_notes.views.sheetView[0].showGridLines = True
    ws_notes["A1"].value = "Key Workload & Architecture Assumptions"
    ws_notes["A1"].font = Font(name="Segoe UI", size=14, bold=True, color="1E3A8A")

    cur_row = 3
    for idx, asm in enumerate(assumptions, 1):
        ws_notes.cell(row=cur_row, column=1, value=f"{idx}. {asm}").font = Font(name="Segoe UI", size=10)
        cur_row += 1

    cur_row += 2
    ws_notes.cell(row=cur_row, column=1, value="Exclusions & Non-Included Cost Drivers").font = Font(name="Segoe UI", size=14, bold=True, color="B91C1C")
    cur_row += 2
    for idx, exc in enumerate(exclusions, 1):
        ws_notes.cell(row=cur_row, column=1, value=f"• {exc}").font = Font(name="Segoe UI", size=10)
        cur_row += 1

    ws_notes.column_dimensions["A"].width = 80

    output = io.BytesIO()
    wb.save(output)
    content_bytes = output.getvalue()

    filename = f"{project_name.replace(' ', '_').lower()}_{provider}_{scenario_name.lower()}_estimate.xlsx"
    return await save_artifact(content_bytes, filename, user_id, job_id)
